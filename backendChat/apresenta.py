import pandas as pd
from pptx import Presentation
from pptx.util import Inches
import os
# Criar uma nova apresentação PowerPoint
prs = Presentation()

# Estrutura da apresentação (título, conteúdo)
slides_content = [
    ("Inteligência Artificial e o Futuro do Trabalho", "Seu Nome\nFaculdade de Engenharia da Universidade do Porto (FEUP)"),
    ("Introdução", "📌 Quem aqui já usou IA hoje?\n📌 Como a IA está mudando o mundo do trabalho?\n📌 O que vamos aprender hoje?"),
    ("O Que é Inteligência Artificial?", "📌 IA permite que máquinas pensem e aprendam como humanos.\n📌 Tipos:\n   - IA Fraca (Siri, Alexa)\n   - IA Forte (Futuro: máquinas que pensam sozinhas)"),
    ("Como a IA Funciona?", "📌 Machine Learning e Redes Neurais\n📌 Exemplo: Como o YouTube recomenda vídeos para você?"),
    ("IA no Nosso Dia a Dia", "📌 Exemplos:\n   - Spotify (músicas)\n   - Netflix (filmes)\n   - TikTok (vídeos)"),
    ("IA e o Mercado de Trabalho", "📌 Profissões que estão mudando:\n   - Carros autônomos 🚗\n   - Atendimento automático 🤖\n   - Medicina com IA 🏥"),
    ("A IA Vai Roubar Todos os Empregos?", "📌 MITO! 🚫 IA cria novas profissões:\n   - Engenheiro de IA 🤖\n   - Cientista de Dados 📊\n   - Especialista em Ética da IA ⚖️"),
    ("Como a FEUP Prepara Você para Esse Futuro?", "📌 Cursos na FEUP:\n   - Engenharia Informática e Computação 💻\n   - Projetos e laboratórios de IA 🧪"),
    ("Demonstração de IA", "📌 Exemplo prático (ChatGPT, reconhecimento de imagens)\n📌 Mostrar um chatbot ou ferramenta de IA"),
    ("Como Aprender IA Desde Já?", "📌 Dicas para começar:\n   - Aprender Python 🐍\n   - Usar Google Colab 🤖\n   - Participar de hackathons"),
    ("Oportunidades e Futuro da IA", "📌 O que esperar do futuro?\n   - IA em robôs humanoides 🤖\n   - Computação quântica ⚛️\n   - IA na exploração espacial 🚀"),
    ("Conclusão e Chamado para Ação", "📌 Resumo:\n   - IA já está no nosso dia a dia\n   - Oportunidade de estudar e se tornar um especialista\n📌 Quem aqui quer trabalhar com IA?\n📌 Explora a FEUP!")
]

# Adicionar slides com conteúdo
for title, content in slides_content:
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # Layout com título e conteúdo
    title_shape = slide.shapes.title
    content_shape = slide.placeholders[1]

    title_shape.text = title
    content_shape.text = content

# Salvar a apresentação
pptx_path = os.getcwd()+ "IA_Futuro_Trabalho_FEUP.pptx"
prs.save(pptx_path)

